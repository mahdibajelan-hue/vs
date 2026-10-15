#!/usr/bin/env python3
"""Structural conformity of the approved grid templates against the ORIGINAL Excel workbooks (independent of the converter's own output):
the rendered, data-free sheet (scripts/print-check.mjs --blank) is compared with the source worksheet — text of every cell, merged ranges, column widths,
row heights, page setup, number of pictures, and the SHA-256 of the workbook the template says it came from.
It does NOT compare pixels (no printed sample is available here) — that stays a human comparison.
usage: python3 -I scripts/conformity-check.py --welding W.xlsx --radio R.xlsx --blank DIR/blank [--out report.md]"""
import argparse, hashlib, json, re, sys, glob, os
from html.parser import HTMLParser
import openpyxl, zipfile
from openpyxl.utils import get_column_letter, range_boundaries

ap = argparse.ArgumentParser(); ap.add_argument('--welding', required=True); ap.add_argument('--radio', required=True); ap.add_argument('--blank', required=True); ap.add_argument('--templates', default='src/modules/execution/forms/templates'); ap.add_argument('--out'); ap.add_argument('--verbose', action='store_true')
a = ap.parse_args()
FILES = {'welding.xlsx': a.welding, 'w.xlsx': a.welding, 'radiographi.xlsx': a.radio, 'r.xlsx': a.radio}
norm = lambda s: re.sub(r'\s+', ' ', (s or '').replace('‌', '')).strip()

class P(HTMLParser):
    def __init__(s): super().__init__(); s.tds = {}; s.cur = None; s.cols = []; s.rows = []; s.in_tr = False
    def handle_starttag(s, t, at):
        d = dict(at)
        if t == 'col': s.cols.append(float(re.search(r'width:([\d.]+)px', d.get('style', '')).group(1)))
        if t == 'tr': s.rows.append(float(re.search(r'height:([\d.]+)px', d.get('style', '')).group(1)))
        if t == 'td': s.cur = d['data-cell']; s.tds[s.cur] = {'cs': int(d.get('colspan', 1)), 'rs': int(d.get('rowspan', 1)), 'text': ''}
    def handle_endtag(s, t):
        if t == 'td': s.cur = None
    def handle_data(s, data):
        if s.cur: s.tds[s.cur]['text'] += data

def colw(ws, idx):
    for d in ws.column_dimensions.values():
        if d.min and d.max and d.min <= idx <= d.max and d.width: return d.width
    return ws.sheet_format.defaultColWidth or 8.43
def zip_pics(path, sheet_title):
    z = zipfile.ZipFile(path); wbx = z.read('xl/workbook.xml').decode(); rel = z.read('xl/_rels/workbook.xml.rels').decode()
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(\w+)"' % re.escape(sheet_title), wbx).group(1)
    tgt = re.search(r'Id="%s"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="%s"' % (rid, rid), rel); tgt = tgt.group(1) or tgt.group(2)
    sp = 'xl/' + tgt.lstrip('/').replace('xl/', '')
    rp = sp.replace('worksheets/', 'worksheets/_rels/') + '.rels'
    if rp not in z.namelist(): return 0
    m = re.search(r'Target="\.\./drawings/(drawing\d+\.xml)"', z.read(rp).decode())
    return len(re.findall(r'<xdr:pic>|<pic>', z.read('xl/drawings/' + m.group(1)).decode())) if m else 0
WB = {}
out = []; total_bad = 0
out.append('# Structural conformity: approved grid templates vs original workbooks\n')
out.append('| form | sheet | source hash | texts | merges | col widths | row heights | page setup | pictures |\n|---|---|---|---|---|---|---|---|---|')
for f in sorted(glob.glob(os.path.join(a.templates, '*.grid.json'))):
    t = json.load(open(f)); key = t['formKey']
    path = FILES[t['source']['file']]
    sha_ok = hashlib.sha256(open(path, 'rb').read()).hexdigest() == t['source']['sha256']
    if path not in WB: WB[path] = (openpyxl.load_workbook(path, data_only=True), openpyxl.load_workbook(path))
    ws = WB[path][0][t['sheet']]; wf = WB[path][1][t['sheet']]
    minc, minr, maxc, maxr = range_boundaries(t['area'])
    p = P(); p.feed(open(os.path.join(a.blank, key + '.html'), encoding='utf8').read())
    hidden = {i for i, c in enumerate(t['cols']) if c['hidden']}
    # texts
    bad_t = 0; n_t = 0; n_formula = [0]
    for row in ws.iter_rows(min_row=minr, max_row=maxr, min_col=minc, max_col=maxc):
        for c in row:
            if c.value is None or norm(str(c.value)) == '': continue
            if c.__class__.__name__ == 'MergedCell': continue
            fv = wf.cell(row=c.row, column=c.column).value
            if isinstance(fv, str) and fv.startswith('='): n_formula[0] += 1; continue
            if (c.column - minc) in hidden: continue
            n_t += 1
            addr = f'{get_column_letter(c.column - minc + 1)}{c.row - minr + 1}'
            got = p.tds.get(addr, {}).get('text', None)
            if got is None or norm(got) != norm(str(c.value)):
                bad_t += 1
                if a.verbose: print('TEXT', key, addr, repr(c.value), repr(got), file=sys.stderr)
    # merges
    bad_m = 0; n_m = 0
    for m in ws.merged_cells.ranges:
        if m.max_row < minr or m.min_row > maxr or m.max_col < minc or m.min_col > maxc: continue
        if all((c - minc) in hidden for c in range(m.min_col, m.max_col + 1)): continue  # merge lives entirely in hidden columns: nothing to print
        n_m += 1
        addr = f'{get_column_letter(m.min_col - minc + 1)}{m.min_row - minr + 1}'
        td = p.tds.get(addr)
        exp_cs = sum(1 for c in range(m.min_col, m.max_col + 1) if (c - minc) not in hidden)
        if td is None or td['rs'] != (m.max_row - m.min_row + 1) or td['cs'] != max(exp_cs, 1):
            bad_m += 1
            if a.verbose: print('MERGE', key, m, td, file=sys.stderr)
    # widths / heights
    visible = [i for i in range(len(t['cols'])) if i not in hidden]
    bad_w = 0
    for k, i in enumerate(visible):
        w = colw(ws, minc + i)
        if abs(p.cols[k] - round(w * 7 + 5)) > 1:
            bad_w += 1
            if a.verbose: print('WIDTH', key, k, p.cols[k], w, file=sys.stderr)
    bad_h = 0
    vis_rows = [i for i, r in enumerate(t['rows']) if not r['hidden']]
    for k, i in enumerate(vis_rows):
        d = ws.row_dimensions[minr + i]; h = d.height if d and d.height else (ws.sheet_format.defaultRowHeight or 15)
        if abs(p.rows[k] - h * 96 / 72) > 1.5: bad_h += 1
    # page setup
    ps = ws.page_setup; pm = ws.page_margins
    ok_ps = (ps.orientation or 'portrait') == t['print']['orientation'] and abs((ps.scale or 100) - (t['print']['scale'] or 100)) < 1 and abs((pm.left or 0) - t['print']['margins']['left']) < 0.01 and abs((pm.top or 0) - t['print']['margins']['top']) < 0.01
    pics = zip_pics(path, t['sheet']); tp = len([l for l in t['logos'] if l['kind'] == 'pic'])
    bad = bad_t + bad_m + bad_w + bad_h + (0 if sha_ok else 1) + (0 if ok_ps else 1) + (0 if pics >= tp else 1); total_bad += bad
    out.append(f"| {key} | {t['sheet']} | {'✓' if sha_ok else '✗'} | {n_t - bad_t}/{n_t} (+{n_formula[0]} formula) | {n_m - bad_m}/{n_m} | {len(visible) - bad_w}/{len(visible)} | {len(vis_rows) - bad_h}/{len(vis_rows)} | {'✓' if ok_ps else '✗'} | {tp} of {pics} |")
out.append(f'\nmismatches in total: {total_bad}')
txt = '\n'.join(out); print(txt)
if a.out: open(a.out, 'w').write(txt)
sys.exit(1 if total_bad else 0)
