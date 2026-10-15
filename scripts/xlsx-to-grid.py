#!/usr/bin/env python3
"""Convert one worksheet of an approved Excel form into a frozen grid template (JSON).

    python3 scripts/xlsx-to-grid.py <file.xlsx> <sheet name> <out.json> [--form-key welding]

The template keeps the geometry of the approved sheet exactly: column widths, row heights, merged ranges, fonts, alignment, fills,
borders, fixed texts, print area / orientation / scale and the anchors of logos. Nothing is re-laid-out. Cell values are NOT interpreted:
the binding of fields to cells lives next to the template (see src/modules/execution/forms/*.bindings.ts) and must be approved per form.
Needs: openpyxl. The source workbook is read as-is (external links are ignored, never followed).
"""
import hashlib, json, re, sys, zipfile
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

THEME = ['FFFFFF', '000000', 'EEECE1', '1F497D', '4F81BD', 'C0504D', '9BBB59', '8064A2', '4BACC6', 'F79646']
INDEXED = {0: '000000', 1: 'FFFFFF', 2: 'FF0000', 3: '00FF00', 4: '0000FF', 5: 'FFFF00', 6: 'FF00FF', 7: '00FFFF', 8: '000000', 9: 'FFFFFF', 22: 'C0C0C0', 23: '808080', 64: '000000'}


def tint(hexrgb, t):
    r, g, b = (int(hexrgb[i:i + 2], 16) for i in (0, 2, 4))
    f = (lambda c: round(c * (1 + t))) if t < 0 else (lambda c: round(c + (255 - c) * t))
    return '%02X%02X%02X' % tuple(max(0, min(255, f(c))) for c in (r, g, b))


def color(c):
    if c is None:
        return None
    try:
        if c.type == 'rgb' and isinstance(c.rgb, str) and len(c.rgb) in (6, 8):
            return c.rgb[-6:]
        if c.type == 'theme' and c.theme is not None and c.theme < len(THEME):
            return tint(THEME[c.theme], c.tint or 0)
        if c.type == 'indexed' and c.indexed in INDEXED:
            return INDEXED[c.indexed]
    except Exception:
        return None
    return None


def side(s):
    if s is None or s.style is None:
        return None
    return [s.style, color(s.color) or '000000']


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    src, sheet, out = sys.argv[1:4]
    key = sys.argv[sys.argv.index('--form-key') + 1] if '--form-key' in sys.argv else sheet.lower().replace(' ', '_')
    wb = openpyxl.load_workbook(src)
    ws = wb[sheet]
    pa = ws.print_area
    area = pa if isinstance(pa, str) else (pa[0] if pa else None)
    area = area.split('!')[-1].replace('$', '') if area else ws.dimensions
    c1, r1, c2, r2 = range_boundaries(area)
    # columns
    widths = {}
    hidden = set()
    for k, d in ws.column_dimensions.items():
        lo, hi = d.min or openpyxl.utils.column_index_from_string(k), d.max or openpyxl.utils.column_index_from_string(k)
        for ci in range(lo, min(hi, c2) + 1):
            if d.width:
                widths[ci] = d.width
            if d.hidden:
                hidden.add(ci)
    dw = ws.sheet_format.defaultColWidth or ws.sheet_format.baseColWidth or 8.43
    cols = [{'w': round((widths.get(ci, dw)) * 7 + 5), 'hidden': ci in hidden} for ci in range(c1, c2 + 1)]
    drh = ws.sheet_format.defaultRowHeight or 15
    rows = []
    for ri in range(r1, r2 + 1):
        d = ws.row_dimensions.get(ri)
        rows.append({'h': round(((d.height if d and d.height else drh)) * 96 / 72), 'hidden': bool(d and d.hidden)})
    merged = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        a, b, c, d = m.min_col, m.min_row, m.max_col, m.max_row
        if c < c1 or a > c2 or d < r1 or b > r2:
            continue
        merged[(b, a)] = (min(d, r2) - b + 1, min(c, c2) - a + 1, m)
        for rr in range(b, min(d, r2) + 1):
            for cc in range(a, min(c, c2) + 1):
                if (rr, cc) != (b, a):
                    covered.add((rr, cc))
    styles, sidx = [], {}

    def style_of(cell, rng):
        f, al, fl = cell.font, cell.alignment, cell.fill
        bd = cell.border
        if rng:  # a merged cell takes each edge from the cell that really carries it
            tl = cell
            top, left = tl.border.top, tl.border.left
            right = ws.cell(rng.min_row, rng.max_col).border.right
            bottom = ws.cell(rng.max_row, rng.min_col).border.bottom
        else:
            top, left, right, bottom = bd.top, bd.left, bd.right, bd.bottom
        st = {
            'f': [f.name, f.sz, bool(f.b), bool(f.i), bool(f.u), color(f.color)],
            'a': [al.horizontal, al.vertical, bool(al.wrap_text), al.textRotation or 0, al.readingOrder or 0],
            'b': [side(top), side(right), side(bottom), side(left)],
            'bg': color(fl.fgColor) if fl and fl.fill_type == 'solid' else None,
            'nf': cell.number_format if cell.number_format != 'General' else None,
        }
        k = json.dumps(st, sort_keys=True)
        if k not in sidx:
            sidx[k] = len(styles)
            styles.append(st)
        return sidx[k]

    cells = []
    for ri in range(r1, r2 + 1):
        for ci in range(c1, c2 + 1):
            if (ri, ci) in covered:
                continue
            cell = ws.cell(ri, ci)
            rs, cs, rng = merged.get((ri, ci), (1, 1, None))
            v = cell.value
            s = style_of(cell, rng)
            blank_style = styles[s]['bg'] is None and not any(styles[s]['b']) and s == 0
            if v in (None, '') and rs == 1 and cs == 1 and not styles[s]['bg'] and not any(styles[s]['b']):
                continue
            item = {'r': ri - r1, 'c': ci - c1, 's': s}
            if rs > 1:
                item['rs'] = rs
            if cs > 1:
                item['cs'] = cs
            if v not in (None, ''):
                item['v'] = str(v) if not str(v).startswith('=') else None
                if str(v).startswith('='):
                    item['formula'] = str(v)
            cells.append(item)
    # logos: read drawing anchors straight from the package (EMF cannot be shown by browsers; the template records where each logo sits)
    logos = []
    try:
        z = zipfile.ZipFile(src)
        rel = {}
        for n in z.namelist():
            if n.startswith('xl/worksheets/_rels/') and n.endswith('.rels'):
                rel[n] = z.read(n).decode('utf8', 'ignore')
        wbx = z.read('xl/workbook.xml').decode('utf8', 'ignore')
        wbrels = z.read('xl/_rels/workbook.xml.rels').decode('utf8', 'ignore')
        rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(sheet), wbx)
        tgt = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid.group(1), wbrels) or re.search(r'Target="([^"]+)"[^>]*Id="%s"' % rid.group(1), wbrels)
        sp = 'xl/' + tgt.group(1).lstrip('/').replace('xl/', '')
        rels = rel.get(sp.replace('worksheets/', 'worksheets/_rels/') + '.rels', '')
        dm = re.search(r'Target="\.\./drawings/(drawing\d+\.xml)"', rels)
        if dm:
            dx = z.read('xl/drawings/' + dm.group(1)).decode('utf8', 'ignore')
            drels = z.read('xl/drawings/_rels/' + dm.group(1) + '.rels').decode('utf8', 'ignore')
            for a in re.findall(r'<xdr:(?:twoCellAnchor|oneCellAnchor)[^>]*>.*?</xdr:(?:twoCellAnchor|oneCellAnchor)>', dx, flags=re.S):
                f = re.search(r'<xdr:from><xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', a, flags=re.S)
                t = re.search(r'<xdr:to><xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', a, flags=re.S)
                emb = re.search(r'r:embed="(rId\d+)"', a)
                media = None
                if emb:
                    m = re.search(r'Id="%s"[^>]*Target="\.\./media/([^"]+)"' % emb.group(1), drels) or re.search(r'Target="\.\./media/([^"]+)"[^>]*Id="%s"' % emb.group(1), drels)
                    media = m.group(1) if m else None
                kind = 'pic' if '<xdr:pic' in a else 'shape'
                if f:
                    logos.append({'kind': kind, 'media': media, 'from': [int(f.group(2)) - (r1 - 1), int(f.group(1)) - (c1 - 1)], 'to': [int(t.group(2)) - (r1 - 1), int(t.group(1)) - (c1 - 1)] if t else None, 'text': ' '.join(re.findall(r'<a:t>([^<]*)</a:t>', a)) or None})
    except Exception as e:  # logos are optional metadata
        logos.append({'kind': 'error', 'text': str(e)})
    ps = ws.page_setup
    tpl = {
        'formKey': key, 'sheet': sheet, 'area': area,
        'source': {'file': src.split('/')[-1], 'sha256': hashlib.sha256(open(src, 'rb').read()).hexdigest()},
        'print': {'orientation': ps.orientation or 'portrait', 'paperSize': ps.paperSize, 'scale': ps.scale, 'fitToPage': bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage), 'margins': {k: round(getattr(ws.page_margins, k), 3) for k in ('left', 'right', 'top', 'bottom')}, 'center': bool(ws.print_options.horizontalCentered)},
        'rtl': bool(ws.sheet_view.rightToLeft),
        'cols': cols, 'rows': rows, 'styles': styles, 'cells': cells, 'logos': logos,
        'stats': {'cells': len(cells), 'merged': len(merged), 'texts': sum(1 for c in cells if c.get('v')), 'hiddenCols': len(hidden)},
    }
    with open(out, 'w', encoding='utf8') as fh:
        json.dump(tpl, fh, ensure_ascii=False, separators=(',', ':'))
    print(out, tpl['stats'], 'logos', len(logos))


main()
